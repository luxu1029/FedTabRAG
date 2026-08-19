#!/usr/bin/env python3
"""Insert the verified Task-2 execution record into the next-stage DOCX."""
from __future__ import annotations

import argparse
import csv
import json
import shutil
from pathlib import Path

from docx import Document


TITLE = "任务 2实际执行与验收结果（2026-08-04）"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--document", type=Path,
                        default=Path("FedTabRAG_下一阶段实验工作安排.docx"))
    parser.add_argument("--backup", type=Path,
                        default=Path("FedTabRAG_下一阶段实验工作安排.before_task2_results_20260804.docx"))
    parser.add_argument("--root", type=Path, default=Path("outputs/diagnosis"))
    return parser.parse_args()


def rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def bucket(data: list[dict[str, str]], name: str) -> dict[str, str]:
    return next(item for item in data if item["bucket"] == name)


def main() -> None:
    args = parse_args()
    document = Document(args.document)
    if any(paragraph.text.strip() == TITLE for paragraph in document.paragraphs):
        raise RuntimeError("Task-2 result section already exists; refusing duplicate insertion")
    anchor = next((paragraph for paragraph in document.paragraphs
                   if paragraph.text.strip().startswith("任务 3：")), None)
    if anchor is None:
        raise RuntimeError("Task-3 anchor not found")
    manifest = json.loads((args.root / "bucket_input_sha256.json").read_text())
    if len(manifest["inputs"]) != 3 or len(manifest["outputs"]) != 15:
        raise ValueError("Task-2 manifest did not pass cardinality audit")
    steds = rows(args.root / "buckets/bucket_by_steds_dev.csv")
    cell_count = rows(args.root / "buckets/bucket_by_cell_count_dev.csv")
    table_size = rows(args.root / "buckets/bucket_by_table_size_dev.csv")
    evidence = rows(args.root / "buckets/bucket_by_evidence_level_dev.csv")
    question = rows(args.root / "buckets/bucket_by_question_type_dev.csv")
    if not args.backup.exists():
        shutil.copy2(args.document, args.backup)

    elements = []

    def add(text: str, style: str | None = None) -> None:
        paragraph = document.add_paragraph(text, style=style)
        elements.append(paragraph._element)

    add(TITLE, "Heading 3")
    add("执行状态：PASS。任务2已完成全部分桶、摘要、日志和SHA256验收。")
    add(
        "实现脚本：src/evaluation/bucket_flat_unitable_diagnosis.py。脚本支持CLI、seed、"
        "日志、显式overwrite、原子写入、ID/结构覆盖检查和桶加权回算。", "List Bullet")
    add(
        "协议冻结：S-TEDS四分位只由dev计算，边界为Q25=0.909091、Q50=0.933333、"
        "Q75=0.947368；同一边界原样应用到test。表格大小使用small≤20、medium≤50、"
        "large>50个单元格。", "List Bullet")
    add(
        "输出包括dev/test各6类分桶（S-TEDS、cell-count、table-size、question-type、"
        "program-type、evidence-level）、2份逐query结构合并表和diagnosis_summary.md。",
        "List Bullet")
    add(
        "结构合并字段完整覆盖S-TEDS、valid_html、cell-count exact/ratio、confidence、"
        "预测/真实cell、row、column数及其绝对误差。", "List Bullet")

    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ("dev分桶", "count", "Flat NDCG@10", "UniTable NDCG@10",
               "delta_ndcg", "UniTable win rate")
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    selected = [
        ("S-TEDS q1_low", bucket(steds, "q1_low")),
        ("S-TEDS q4_high", bucket(steds, "q4_high")),
        ("cell-count exact", bucket(cell_count, "exact")),
        ("cell-count mismatch", bucket(cell_count, "mismatch")),
        ("table-size medium", bucket(table_size, "medium")),
        ("evidence cell", bucket(evidence, "cell")),
        ("question hybrid", bucket(question, "hybrid")),
    ]
    for label, item in selected:
        values = (label, item["count"], f"{float(item['flat_ndcg@10']):.6f}",
                  f"{float(item['unitable_ndcg@10']):.6f}",
                  f"{float(item['delta_ndcg']):.6f}",
                  f"{float(item['unitable_win_rate']):.4f}")
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    elements.append(table._element)

    add(
        "dev总体mean delta_ndcg=-0.003463；test总体为+0.000107，但test仅作描述，"
        "不用于选择任何阈值、融合权重或方案。")
    add(
        "关键诊断：高S-TEDS的q4_high仍下降0.005826，且下降大于q1_low的0.003853，"
        "说明仅提高结构识别质量不足以保证检索收益，当前证据更倾向于结构序列化/编码问题。")
    add(
        "cell-count mismatch只有11条，delta=-0.005784，差于exact桶的-0.003434；"
        "它可以作为后续门控候选特征，但样本很少，任务2不能据此直接冻结阈值。")
    add(
        "各表格大小桶均为负增益，medium桶退化最小（-0.001898）；cell证据层级"
        "delta=-0.003615，table层级delta=-0.002937。hybrid问题出现+0.000631的局部"
        "正值，但不能代表稳定总体收益，仍需任务4的dev门控实验验证。")
    add("任务2最终验收：", "Heading 3")
    add("☑ dev 883条、test 1147条结构字段完整合并，无缺失、无重复。")
    add("☑ 每个分桶均包含规定指标、count和win/both-fail rate。")
    add("☑ 六个维度的桶count均覆盖全集，按count加权可精确回算总体指标。")
    add("☑ 3个输入和15个输出的SHA256全部复算一致。")
    add("☑ test未参与分桶边界或方案选择；任务2只做诊断，不冻结门控规则。")
    add("☑ diagnosis_summary.md已指出高S-TEDS仍退化，可进入任务3及后续dev实验。")

    for element in elements:
        anchor._p.addprevious(element)
    document.save(args.document)
    print(f"updated={args.document}")
    print(f"backup={args.backup}")


if __name__ == "__main__":
    main()
