#!/usr/bin/env python3
"""Insert the verified Task-3 case-export record into the next-stage DOCX."""
from __future__ import annotations

import argparse
import csv
import json
import shutil
from pathlib import Path

from docx import Document


TITLE = "任务 3实际执行与验收结果（2026-08-04）"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--document", type=Path,
                        default=Path("FedTabRAG_下一阶段实验工作安排.docx"))
    parser.add_argument("--backup", type=Path,
                        default=Path("FedTabRAG_下一阶段实验工作安排.before_task3_results_20260804.docx"))
    parser.add_argument("--root", type=Path, default=Path("outputs/diagnosis"))
    return parser.parse_args()


def selected_rows(path: Path) -> dict[str, list[dict]]:
    with path.open(newline="", encoding="utf-8") as handle:
        source = list(csv.DictReader(handle))
    return {
        "flat_win": sorted(
            (row for row in source if row["winner"] == "flat_win"),
            key=lambda row: (float(row["delta_ndcg"]), row["query_id"]))[:30],
        "unitable_win": sorted(
            (row for row in source if row["winner"] == "unitable_win"),
            key=lambda row: (-float(row["delta_ndcg"]), row["query_id"]))[:30],
        "both_fail": sorted(
            (row for row in source if row["winner"] == "both_fail"),
            key=lambda row: (
                float(row["flat_ndcg@10"]) + float(row["unitable_ndcg@10"]),
                -min(int(row["flat_gold_rank"]), int(row["unitable_gold_rank"])),
                row["query_id"]),
        )[:30],
    }


def main() -> None:
    args = parse_args()
    document = Document(args.document)
    if any(paragraph.text.strip() == TITLE for paragraph in document.paragraphs):
        raise RuntimeError("Task-3 result section already exists; refusing duplicate insertion")
    anchor = next((paragraph for paragraph in document.paragraphs
                   if paragraph.text.strip().startswith("任务 4：")), None)
    if anchor is None:
        raise RuntimeError("Task-4 anchor not found")
    manifest = json.loads((args.root / "manual_cases_input_sha256.json").read_text())
    if len(manifest["inputs"]) != 7 or len(manifest["outputs"]) != 3:
        raise ValueError("Task-3 SHA manifest did not pass cardinality audit")
    cases = selected_rows(args.root / "pairwise_flat_unitable_test.csv")
    if any(len(rows) != 30 for rows in cases.values()):
        raise ValueError("Task-3 workbook row count mismatch")
    for name in cases:
        if not (args.root / f"manual_cases_{name}.xlsx").exists():
            raise FileNotFoundError(args.root / f"manual_cases_{name}.xlsx")
    if not args.backup.exists():
        shutil.copy2(args.document, args.backup)

    elements = []

    def add(text: str, style: str | None = None) -> None:
        paragraph = document.add_paragraph(text, style=style)
        elements.append(paragraph._element)

    add(TITLE, "Heading 3")
    add("执行状态：案例导出与技术验收PASS；人工错误归因尚待填写，不生成自动结论。")
    add(
        "实现脚本：src/evaluation/export_manual_cases.py。脚本使用固定规则从test配对表"
        "导出案例，支持CLI、seed、日志、显式overwrite、输入一致性检查和SHA256。",
        "List Bullet")
    add(
        "输出：outputs/diagnosis/manual_cases_flat_win.xlsx、"
        "manual_cases_unitable_win.xlsx、manual_cases_both_fail.xlsx；每份30条。",
        "List Bullet")
    add(
        "每条案例均包含问题与元数据、两路gold evidence、两路Top-5证据及分数、"
        "gold rank、S-TEDS、cell-count consistency、预测/真实行列单元格数和"
        "UniTable预测HTML摘要。", "List Bullet")
    add(
        "error_type、manual_note、whether_structure_helpful三列保持空白，并配置下拉候选；"
        "脚本没有自动填写归因，也没有据test案例选择阈值或调整模型。", "List Bullet")

    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ("案例组", "数量", "固定选择规则", "delta最小值", "delta最大值")
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    rules = {
        "flat_win": "delta升序，query_id打破并列",
        "unitable_win": "delta降序，query_id打破并列",
        "both_fail": "双路Top-10未命中；NDCG升序、最佳rank降序",
    }
    labels = {"flat_win": "Flat胜", "unitable_win": "UniTable胜",
              "both_fail": "二者都失败"}
    for name in ("flat_win", "unitable_win", "both_fail"):
        values = [float(row["delta_ndcg"]) for row in cases[name]]
        row_values = (labels[name], str(len(values)), rules[name],
                      f"{min(values):.6f}", f"{max(values):.6f}")
        for cell, value in zip(table.add_row().cells, row_values):
            cell.text = value
    elements.append(table._element)

    add(
        "技术验收结果：90个query互不重复；Flat胜delta范围[-0.369070,-0.130930]；"
        "UniTable胜范围[0.130127,0.410392]；both_fail的两路NDCG均为0且两路gold rank"
        "均大于10。")
    add(
        "SHA256验收：7个输入和3个Excel输出全部复算一致。输入包括配对表、两路冻结"
        "检索结果、两路语料、结构指标和UniTable原始结构预测。")
    add("任务3当前验收：", "Heading 3")
    add("☑ 三类案例各30条，共90条；抽取顺序固定，无人工挑选。")
    add("☑ question、gold、Flat/UniTable Top-5和全部结构摘要字段完整。")
    add("☑ Excel可直接人工填写，三个人工归因列初始为空。")
    add("☑ 未修改配对表、检索结果、结构预测或模型。")
    add("☐ 人工研究者尚未完成90条error_type、manual_note和structure_helpful标注。")
    add("☐ manual_case_analysis_summary.md须在人工标注完成后撰写，当前不得自动下结论。")

    for element in elements:
        anchor._p.addprevious(element)
    document.save(args.document)
    print(f"updated={args.document}")
    print(f"backup={args.backup}")


if __name__ == "__main__":
    main()
