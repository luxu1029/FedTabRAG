#!/usr/bin/env python3
"""Insert frozen Day-10 results and next-step feasibility analysis into the tutorial."""
from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK


TITLE = "Day 10实际执行结果与下一步可行性分析（2026-08-03）"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--document", type=Path,
                        default=Path("FedTabRAG_10天服务器实验详细执行教程.docx"))
    parser.add_argument("--backup", type=Path,
                        default=Path("FedTabRAG_10天服务器实验详细执行教程.before_day10_results_20260803.docx"))
    return parser.parse_args()


def add_table(document: Document, headers: list[str], rows: list[list[str]],
              elements: list) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    elements.append(table._element)


def main() -> None:
    args = parse_args()
    if not args.document.exists():
        raise FileNotFoundError(args.document)
    document = Document(args.document)
    if any(paragraph.text.strip() == TITLE for paragraph in document.paragraphs):
        raise RuntimeError("Day-10 result section already exists; refusing duplicate insertion")
    anchor = next((paragraph for paragraph in document.paragraphs
                   if paragraph.text.strip().startswith("附录A｜")), None)
    if anchor is None:
        raise RuntimeError("Appendix-A anchor not found")
    if not args.backup.exists():
        shutil.copy2(args.document, args.backup)

    elements = []

    def paragraph(text: str = "", style: str | None = None):
        item = document.add_paragraph(text, style=style)
        elements.append(item._element)
        return item

    page = paragraph()
    page.add_run().add_break(WD_BREAK.PAGE)
    paragraph(TITLE, "Heading 1")
    paragraph("一、冻结状态与验收结论", "Heading 2")
    paragraph(
        "Day 10已完成并通过最终验收。收敛图直接读取原始convergence.csv；"
        "结构扰动只在预测结构语料副本中执行；所有表格均采用strict missing策略，"
        "缺失项写为NA，不以假设值或回填值替代。seed=42的既有结果保持冻结，"
        "并在资源允许条件下补完seed=7和2026。"
    )
    paragraph("冻结产物路径：outputs/day10/。", "List Bullet")
    paragraph("收敛图：outputs/day10/figures/three_method_convergence.png（另有PDF）。",
              "List Bullet")
    paragraph("结果表：outputs/day10/tables/；共8张CSV表。", "List Bullet")
    paragraph("SHA256清单：outputs/day10/sha256_manifest.json。", "List Bullet")
    paragraph("无权重冻结包：outputs/day10/FedTabRAG_day10_frozen_results.tar.gz。",
              "List Bullet")
    paragraph(
        "冻结包SHA256：2eaeae747a20df3a1f720c626405083dc9ed3ccddf21b91a1c9afe0b358f6656；"
        "归档包含68个成员，模型权重文件数为0。", "List Bullet")
    paragraph(
        "当前工作区不是Git仓库，因此git_commit严格记为NA并写明原因；"
        "不得伪造commit。配置、结果和关键脚本均已逐文件保存SHA256。", "List Bullet")

    paragraph("二、主结果与多种子稳定性", "Heading 2")
    add_table(document,
              ["方法", "Table NDCG@10", "Row NDCG@10", "Cell NDCG@10",
               "加权NDCG@10", "100题Numerical EM"],
              [
                  ["FedE4RAG-Flat", "0.314504", "0.442569", "0.120965",
                   "0.256154", "0.07"],
                  ["FedE4RAG-UniTable", "0.310078", "0.329254", "0.138244",
                   "0.229914", "0.07"],
                  ["FedTabRAG-Full（R+U）", "0.341971", "0.307879", "0.131527",
                   "0.226521", "0.05"],
              ], elements)
    paragraph(
        "注：加权NDCG@10采用Table/Row/Cell=0.2/0.3/0.5。问答方法没有输出program，"
        "因此Program/Execution Accuracy记为NA，只报告实际可计算的Numerical EM。"
    )
    add_table(document,
              ["方法", "seed 42", "seed 7", "seed 2026", "三seed均值", "总体标准差"],
              [
                  ["FedE4RAG-Flat", "0.256154", "0.255386", "0.250019",
                   "0.253853", "0.002729"],
                  ["FedE4RAG-UniTable", "0.229914", "0.225312", "0.229439",
                   "0.228222", "0.002067"],
                  ["FedTabRAG-Full（R+U）", "0.226521", "0.226989", "0.226732",
                   "0.226747", "0.000191"],
              ], elements)
    paragraph(
        "补种子结果没有改变排序。Flat稳定领先；UniTable较Flat均值低约10.1%，"
        "完整R+U较Flat均值低约10.7%。R+U的三个seed方差很小，说明当前差距并非"
        "单次随机波动。seed=7和2026的R+U均由dev weighted NDCG选择第4轮best，"
        "冻结后各执行一次test。"
    )

    paragraph("三、结构识别与结构扰动鲁棒性", "Heading 2")
    add_table(document,
              ["方法", "S-TEDS", "Valid HTML", "Cell-count consistency"],
              [
                  ["Flat", "NA", "NA", "NA"],
                  ["GT-oracle", "NA", "NA", "NA"],
                  ["UniTable（test）", "0.922466", "1.000000", "0.977332"],
              ], elements)
    paragraph(
        "Flat没有结构识别输出；GT没有独立测量文件，因此两者均保持NA，"
        "不把GT按假设填成1.0。UniTable结构识别质量较高，但该质量尚未转化为"
        "相对Flat的检索收益。"
    )
    add_table(document,
              ["结构扰动率", "加权NDCG@10", "相对0%绝对下降", "文本多重集合一致"],
              [
                  ["0%", "0.226521", "0.000000", "原始语料"],
                  ["5%", "0.224342", "0.002179", "是"],
                  ["10%", "0.222680", "0.003841", "是"],
                  ["20%", "0.221926", "0.004595", "是"],
                  ["30%", "0.217574", "0.008947", "是"],
              ], elements)
    paragraph(
        "四档扰动分别作用于57/115/229/344张表，全部成功，失败和跳过数均为0。"
        "证据ID、cell_text字段以及全局28205个单元格文本的多重集合保持不变。"
        "随着扰动率上升，检索指标总体下降，证明检索器确实使用了结构；但当前"
        "预测结构带来的有效信号不足以抵消序列化误差和槽位错配噪声。"
    )

    paragraph("四、strict missing审计与可解释结论", "Heading 2")
    paragraph("最终strict missing共9项，均明确写为NA：", None)
    paragraph("Flat和UniTable原始日志未记录逐轮dev weighted NDCG。", "List Bullet")
    paragraph("三种方法原始日志均未记录逐轮Cell Recall@5。", "List Bullet")
    paragraph("Flat与GT没有可追溯的结构识别测量结果。", "List Bullet")
    paragraph("工作区没有Git commit；所有问答方法没有预测program。", "List Bullet")
    paragraph(
        "可写入论文的结论：在FinQA受控渲染表格、oracle单元格文本、"
        "bge-base-en-v1.5和当前联邦预算下，Flat文本检索仍是最强且最稳定的基线。"
        "UniTable具有较高结构指标，但预测结构序列化没有改善下游检索；Hierarchy、"
        "Hard、KD、本地队列与R+U聚合也未弥补该差距。这是一项可复现的负结果，"
        "不应表述为真实PDF端到端OCR系统已经优于文本基线。"
    )

    paragraph("五、下一步可行性分析与执行顺序", "Heading 2")
    paragraph(
        "下一步不宜立即扩大联邦训练预算或继续叠加损失。应先定位Flat与UniTable"
        "约0.026加权NDCG差距的来源，再决定是否进行新一轮训练。"
    )
    paragraph("步骤1：结构质量—检索损失配对诊断", "Heading 3")
    paragraph(
        "按S-TEDS、cell-count consistency、表格大小和结构错误类型分桶；"
        "对同一问题逐条比较Flat与UniTable的gold evidence rank、Recall@5和NDCG@10。"
        "必须分别定位Table、Row、Cell层级的收益与损失，并区分结构识别误差、"
        "序列化误差和检索器表示退化。"
    )
    paragraph(
        "判定：若高S-TEDS桶仍显著落后Flat，则主要瓶颈是结构序列化/编码方式；"
        "若差距集中于低S-TEDS或cell-count不一致样本，则优先做置信度门控。"
    )
    paragraph("步骤2：冻结BGE的最小结构输入对照", "Heading 3")
    paragraph(
        "保持同一问题、候选集、row-major单元格文本集合和冻结BGE，仅比较Flat、GT、"
        "UniTable、UniTable置信度门控、Flat/UniTable分数后融合。所有门控阈值和"
        "融合权重只能由dev确定；test在方案冻结后只评测一次。"
    )
    paragraph("预注册进入联邦训练门槛：", None)
    paragraph("dev加权NDCG@10至少超过冻结Flat 0.003。", "List Bullet")
    paragraph("Row或Cell Recall@5至少提升1个百分点。", "List Bullet")
    paragraph("worst-client相对Flat下降不得超过0.005。", "List Bullet")
    paragraph("至少两个seed方向一致，且提升不能只来自单个结构质量桶。", "List Bullet")
    paragraph("步骤3：优先验证置信度门控与分数后融合", "Heading 3")
    paragraph(
        "高置信度表使用UniTable结构，低置信度表回退Flat；或对两路检索分数做固定"
        "dev权重后融合。该方向不重新训练UniTable，也不改变单元格文本集合，"
        "可直接验证结构是否只在部分样本上有净正贡献。"
    )
    paragraph("步骤4：仅在门槛通过后进行最小联邦验证", "Heading 3")
    paragraph(
        "先跑2客户端3轮smoke，再运行5客户端、seed=42/7/2026。固定company映射、"
        "FedAvg、LoRA、top-k、候选队列和训练预算；按dev weighted NDCG选模，"
        "冻结best后test只运行一次。不得同时修改门控、损失和聚合，以免无法归因。"
    )
    paragraph("步骤5：停止规则与备选方向", "Heading 3")
    paragraph(
        "若门控和后融合仍未达到预注册门槛，停止继续优化当前结构联邦分支，"
        "以Flat作为最终主基线，将UniTable/R+U作为负结果和鲁棒性研究报告。"
        "后续资源转向问答端的数值推理、证据组合或program生成，而不是继续试探"
        "学习率或叠加新的检索损失。"
    )

    paragraph("六、下一阶段验收清单", "Heading 2")
    paragraph("□ 分桶和配对分析覆盖冻结test全集，不挑题、不改题。")
    paragraph("□ 所有输入分支使用完全相同的单元格文本多重集合与候选集合。")
    paragraph("□ 门控阈值/融合权重只由dev确定，test不参与方案选择。")
    paragraph("□ 达到预注册门槛才启动新联邦训练；未达到则保留负结果并停止扩展。")
    paragraph("□ 新结果继续执行strict missing、单次test、SHA256和无权重冻结。")

    for element in elements:
        anchor._p.addprevious(element)
    document.save(args.document)
    print(f"updated={args.document}")
    print(f"backup={args.backup}")


if __name__ == "__main__":
    main()
