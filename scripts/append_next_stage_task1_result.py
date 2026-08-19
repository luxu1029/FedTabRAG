#!/usr/bin/env python3
"""Insert the verified Task-1 execution record into the next-stage DOCX."""
from __future__ import annotations

import argparse
import csv
import json
import shutil
from collections import Counter
from pathlib import Path

from docx import Document


TITLE = "任务 1实际执行与验收结果（2026-08-04）"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--document", type=Path,
                        default=Path("FedTabRAG_下一阶段实验工作安排.docx"))
    parser.add_argument("--backup", type=Path,
                        default=Path("FedTabRAG_下一阶段实验工作安排.before_task1_results_20260804.docx"))
    parser.add_argument("--result-root", type=Path, default=Path("outputs/diagnosis"))
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    document = Document(args.document)
    if any(paragraph.text.strip() == TITLE for paragraph in document.paragraphs):
        raise RuntimeError("Task-1 result section already exists; refusing duplicate insertion")
    anchor = next((paragraph for paragraph in document.paragraphs
                   if paragraph.text.strip().startswith("任务 2：")), None)
    if anchor is None:
        raise RuntimeError("Task-2 anchor not found")

    split_rows = {}
    for split, expected in (("dev", 883), ("test", 1147)):
        path = args.result_root / f"pairwise_flat_unitable_{split}.csv"
        rows = list(csv.DictReader(path.open(encoding="utf-8")))
        if len(rows) != expected or len({row["query_id"] for row in rows}) != expected:
            raise ValueError(f"Unexpected {split} coverage")
        split_rows[split] = rows
    manifest_path = args.result_root / "pairwise_input_sha256.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if len(manifest.get("inputs", [])) != 11 or len(manifest.get("outputs", [])) != 2:
        raise ValueError("Unexpected SHA256 manifest cardinality")

    if not args.backup.exists():
        shutil.copy2(args.document, args.backup)

    elements = []

    def add_paragraph(text: str, style: str | None = None) -> None:
        paragraph = document.add_paragraph(text, style=style)
        elements.append(paragraph._element)

    add_paragraph(TITLE, "Heading 3")
    add_paragraph("执行状态：PASS。任务1已完成，全部验收项通过。")
    add_paragraph(
        "实现脚本：src/evaluation/compare_flat_unitable_pairwise.py。脚本支持"
        "--split dev/test/all、seed、日志、自定义输入输出和显式--overwrite；"
        "包含跨分支query字段、检索协议、gold映射、指标字段、重复ID和覆盖率检查。",
        "List Bullet")
    add_paragraph(
        "现有完整逐query零样本检索结果原先只覆盖test。为形成任务要求的dev配对表，"
        "本次使用冻结BGE和既有只读索引补算Flat与UniTable dev检索结果；"
        "该过程仅做推理，没有训练、重新编码语料或修改模型权重。", "List Bullet")
    add_paragraph(
        "输出：outputs/diagnosis/pairwise_flat_unitable_dev.csv、"
        "pairwise_flat_unitable_test.csv、pairwise_input_sha256.json；"
        "日志：outputs/logs/diagnosis_pairwise.log。", "List Bullet")

    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ("split", "query数", "Flat胜", "UniTable胜", "tie", "both_fail", "覆盖验收")
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for split in ("dev", "test"):
        rows = split_rows[split]
        counts = Counter(row["winner"] for row in rows)
        values = (split, str(len(rows)), str(counts["flat_win"]),
                  str(counts["unitable_win"]), str(counts["tie"]),
                  str(counts["both_fail"]), "无缺失、无重复")
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    elements.append(table._element)

    add_paragraph(
        "winner定义已冻结：两路最佳正例rank都大于10时标记both_fail；否则先比较"
        "NDCG@10，在NDCG相同时比较最佳正例rank，仍相同才标记tie。由于使用全语料"
        "精确排名，所有共享eligible样本均存在大于等于1的gold rank，不需要以-1/inf"
        "表示缺失；both_fail专指两路均未进入Top-10。")
    add_paragraph(
        "单行evidence_level选择两分支共同可评测的最细粒度（cell > row > table）；"
        "gold_evidence_id以JSON对象保存Flat和UniTable各自的多正例ID数组，避免把"
        "预测槽位ID误当作GT槽位ID。")
    add_paragraph(
        "SHA256验收：11个输入和2个输出全部复算一致。dev覆盖883/883，test覆盖"
        "1147/1147；21个规定字段完整，无空值，delta_ndcg重算一致。")
    add_paragraph(
        "边界说明：test在任务1中仅用于固定配对结果和后续人工解释，不用于选择阈值、"
        "融合权重或训练配置。任务4-6仍只能使用dev进行方案选择。")
    add_paragraph("任务1最终验收：", "Heading 3")
    add_paragraph("☑ dev和test每个query均有唯一配对记录。")
    add_paragraph("☑ Flat与UniTable使用相同query、问题文本、gold映射和检索协议。")
    add_paragraph("☑ winner四分类和多正例最细层级规则已明确并通过数值测试。")
    add_paragraph("☑ 未重新训练模型，未修改原始检索结果、语料或索引。")
    add_paragraph("☑ 日志、输入/输出SHA256及固定路径均已保存，可进入任务2。")

    for element in elements:
        anchor._p.addprevious(element)
    document.save(args.document)
    print(f"updated={args.document}")
    print(f"backup={args.backup}")


if __name__ == "__main__":
    main()
