#!/usr/bin/env python3
"""Deterministically export Flat-win, UniTable-win, and both-fail cases for manual review."""
from __future__ import annotations

import argparse
import hashlib
import json
import logging
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


LOG = logging.getLogger("export_manual_cases")
ERROR_TYPES = (
    "structure_prediction_error", "serialization_error", "evidence_mapping_error",
    "retriever_semantic_error", "row_cell_granularity_error", "numeric_reasoning_related",
    "other",
)
FIELDS = (
    "query_id", "sample_id", "question", "company", "question_type", "program_type",
    "evidence_level", "gold_evidence", "flat_top5_evidence", "unitable_top5_evidence",
    "flat_gold_rank", "unitable_gold_rank", "flat_ndcg@10", "unitable_ndcg@10",
    "delta_ndcg", "s_teds", "valid_html", "cell_count_consistency",
    "predicted_cells", "true_cells", "predicted_rows", "true_rows",
    "predicted_cols", "true_cols", "predicted_structure_summary",
    "error_type", "manual_note", "whether_structure_helpful",
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pairwise", type=Path,
                        default=Path("outputs/diagnosis/pairwise_flat_unitable_test.csv"))
    parser.add_argument("--flat-retrieval", type=Path,
                        default=Path("outputs/retrieval/zero_shot_flat.json"))
    parser.add_argument("--unitable-retrieval", type=Path,
                        default=Path("outputs/retrieval/zero_shot_predicted_structure.json"))
    parser.add_argument("--flat-corpus", type=Path,
                        default=Path("data/processed/flat/corpus.jsonl"))
    parser.add_argument("--unitable-corpus", type=Path,
                        default=Path("data/processed/predicted_structure/corpus.jsonl"))
    parser.add_argument("--structure-metrics", type=Path,
                        default=Path("outputs/unitable/structure_metrics_per_sample.jsonl"))
    parser.add_argument("--structure-predictions", type=Path,
                        default=Path("data/interim/unitable_predictions/predictions.jsonl"))
    parser.add_argument("--output-dir", type=Path, default=Path("outputs/diagnosis"))
    parser.add_argument("--sha-output", type=Path,
                        default=Path("outputs/diagnosis/manual_cases_input_sha256.json"))
    parser.add_argument("--log-file", type=Path,
                        default=Path("outputs/logs/diagnosis_manual_cases.log"))
    parser.add_argument("--per-group", type=int, default=30)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--overwrite", action="store_true")
    return parser.parse_args()


def configure_logging(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    formatter = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
    file_handler = logging.FileHandler(path, mode="a", encoding="utf-8")
    stream_handler = logging.StreamHandler()
    file_handler.setFormatter(formatter)
    stream_handler.setFormatter(formatter)
    logging.basicConfig(level=logging.INFO, handlers=[file_handler, stream_handler])


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows = []
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            try:
                rows.append(json.loads(line))
            except json.JSONDecodeError as exc:
                raise ValueError(f"Invalid JSONL {path}:{line_number}: {exc}") from exc
    return rows


def read_csv(path: Path) -> list[dict[str, str]]:
    import csv
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def unique_by(rows: list[dict[str, Any]], key: str, source: Path) -> dict[str, dict[str, Any]]:
    result = {}
    for row in rows:
        value = str(row[key])
        if value in result:
            raise ValueError(f"Duplicate {key}={value} in {source}")
        result[value] = row
    return result


def selected_cases(rows: list[dict[str, str]], count: int) -> dict[str, list[dict[str, str]]]:
    groups = {
        "flat_win": sorted(
            (row for row in rows if row["winner"] == "flat_win"),
            key=lambda row: (float(row["delta_ndcg"]), row["query_id"])),
        "unitable_win": sorted(
            (row for row in rows if row["winner"] == "unitable_win"),
            key=lambda row: (-float(row["delta_ndcg"]), row["query_id"])),
        "both_fail": sorted(
            (row for row in rows if row["winner"] == "both_fail"),
            key=lambda row: (
                float(row["flat_ndcg@10"]) + float(row["unitable_ndcg@10"]),
                -min(int(row["flat_gold_rank"]), int(row["unitable_gold_rank"])),
                row["query_id"],
            )),
    }
    for name, values in groups.items():
        if len(values) < count:
            raise ValueError(f"Need {count} {name} cases, found {len(values)}")
        groups[name] = values[:count]
    selected_ids = [row["query_id"] for values in groups.values() for row in values]
    if len(selected_ids) != len(set(selected_ids)):
        raise AssertionError("Case groups overlap")
    return groups


def retrieval_index(path: Path) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    report = json.loads(path.read_text(encoding="utf-8"))
    return unique_by(report["per_query"], "query_id", path), report.get("protocol", {})


def collect_corpus(path: Path, wanted: set[str]) -> dict[str, dict[str, Any]]:
    found = {}
    with path.open(encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, 1):
            if not line.strip():
                continue
            item = json.loads(line)
            evidence_id = str(item["evidence_id"])
            if evidence_id in wanted:
                if evidence_id in found:
                    raise ValueError(f"Duplicate evidence_id={evidence_id} in {path}:{line_number}")
                found[evidence_id] = item
    missing = sorted(wanted - found.keys())
    if missing:
        raise KeyError(f"Missing {len(missing)} evidence IDs in {path}; first={missing[0]}")
    return found


def compact_evidence(items: list[dict[str, Any]], corpus: dict[str, dict[str, Any]]) -> str:
    values = []
    for item in items:
        evidence_id = str(item["evidence_id"])
        value = {"evidence_id": evidence_id, "text": corpus[evidence_id]["text"]}
        if "rank" in item:
            value["rank"] = int(item["rank"])
        if "score" in item:
            value["score"] = float(item["score"])
        values.append(value)
    text = json.dumps(values, ensure_ascii=False, indent=2)
    if len(text) > 32767:
        raise ValueError("Evidence payload exceeds Excel cell limit; refusing silent truncation")
    return text


def write_xlsx(path: Path, group: str, selection_rule: str, rows: list[dict[str, Any]],
               overwrite: bool, seed: int) -> None:
    if path.exists() and not overwrite:
        raise FileExistsError(f"Refusing to overwrite {path}; use --overwrite")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "cases"
    sheet.append(list(FIELDS))
    for row in rows:
        oversized = [field for field in FIELDS if len(str(row[field])) > 32767]
        if oversized:
            raise ValueError(f"Excel cell limit exceeded for {row['query_id']}: {oversized}")
        sheet.append([row[field] for field in FIELDS])
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {
        "A": 38, "B": 30, "C": 55, "D": 12, "E": 16, "F": 16, "G": 14,
        "H": 70, "I": 70, "J": 70, "Y": 55, "Z": 30, "AA": 55, "AB": 24,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    error_validation = DataValidation(type="list", formula1='"' + ",".join(ERROR_TYPES) + '"')
    helpful_validation = DataValidation(type="list", formula1='"yes,no,uncertain"')
    sheet.add_data_validation(error_validation)
    sheet.add_data_validation(helpful_validation)
    error_validation.add(f"Z2:Z{len(rows)+1}")
    helpful_validation.add(f"AB2:AB{len(rows)+1}")
    metadata = workbook.create_sheet("selection_metadata")
    metadata.append(["field", "value"])
    metadata.append(["group", group])
    metadata.append(["selection_rule", selection_rule])
    metadata.append(["count", len(rows)])
    metadata.append(["seed", seed])
    metadata.append(["automatic_error_conclusion", False])
    metadata.append(["manual_columns_initially_blank", True])
    temporary = path.with_suffix(path.suffix + ".tmp")
    workbook.save(temporary)
    temporary.replace(path)


def main() -> int:
    args = parse_args()
    if args.per_group < 1:
        raise ValueError("--per-group must be positive")
    configure_logging(args.log_file)
    LOG.info("Starting deterministic manual-case export seed=%d per_group=%d",
             args.seed, args.per_group)
    pairwise = read_csv(args.pairwise)
    if any(not row["query_id"].startswith("test:") for row in pairwise):
        raise ValueError("Pairwise input is not test-only")
    groups = selected_cases(pairwise, args.per_group)
    selected = {row["query_id"]: row for values in groups.values() for row in values}
    flat_retrieval, flat_protocol = retrieval_index(args.flat_retrieval)
    unitable_retrieval, unitable_protocol = retrieval_index(args.unitable_retrieval)
    comparable_fields = ("model_fingerprint", "pooling", "normalize", "max_length",
                         "query_instruction", "ranking", "map_definition")
    if any(flat_protocol.get(field) != unitable_protocol.get(field)
           for field in comparable_fields):
        raise ValueError("Flat/UniTable retrieval protocol mismatch")
    structure = unique_by(
        [row for row in read_jsonl(args.structure_metrics) if row.get("split") == "test"],
        "sample_key", args.structure_metrics)
    predictions = unique_by(
        [row for row in read_jsonl(args.structure_predictions) if row.get("split") == "test"],
        "sample_key", args.structure_predictions)
    for source_name, values in (("flat retrieval", flat_retrieval),
                                ("unitable retrieval", unitable_retrieval),
                                ("structure metrics", structure),
                                ("structure predictions", predictions)):
        missing = sorted(set(selected) - values.keys())
        if missing:
            raise KeyError(f"Missing {source_name} for {len(missing)} cases; first={missing[0]}")

    flat_wanted, unitable_wanted = set(), set()
    top_items = {}
    for query_id, row in selected.items():
        level = row["evidence_level"]
        gold = json.loads(row["gold_evidence_id"])
        flat_top = flat_retrieval[query_id]["levels"][level]["top10"][:5]
        unitable_top = unitable_retrieval[query_id]["levels"][level]["top10"][:5]
        if len(flat_top) != 5 or len(unitable_top) != 5:
            raise ValueError(f"Top-5 unavailable for {query_id}/{level}")
        flat_wanted.update(gold["flat"])
        flat_wanted.update(item["evidence_id"] for item in flat_top)
        unitable_wanted.update(gold["unitable"])
        unitable_wanted.update(item["evidence_id"] for item in unitable_top)
        top_items[query_id] = (gold, flat_top, unitable_top)
    flat_corpus = collect_corpus(args.flat_corpus, flat_wanted)
    unitable_corpus = collect_corpus(args.unitable_corpus, unitable_wanted)

    output_specs = {
        "flat_win": ("manual_cases_flat_win.xlsx",
                     "winner=flat_win; delta_ndcg ascending; query_id ascending tie-break"),
        "unitable_win": ("manual_cases_unitable_win.xlsx",
                         "winner=unitable_win; delta_ndcg descending; query_id ascending tie-break"),
        "both_fail": ("manual_cases_both_fail.xlsx",
                      "winner=both_fail; summed NDCG ascending, then best rank descending, query_id"),
    }
    outputs = []
    for group, values in groups.items():
        export_rows = []
        for row in values:
            query_id, level = row["query_id"], row["evidence_level"]
            metric, prediction = structure[query_id], predictions[query_id]
            gold, flat_top, unitable_top = top_items[query_id]
            summary = {
                "status": metric["status"], "valid_html": metric["valid_html"],
                "s_teds": metric["s_teds"], "confidence": metric["confidence"],
                "predicted_shape": [metric["predicted_rows"], metric["predicted_cols"]],
                "true_shape": [metric["true_rows"], metric["true_cols"]],
                "predicted_cells": metric["predicted_cells"],
                "true_cells": metric["true_cells"], "html": prediction.get("html", ""),
            }
            export_rows.append({
                "query_id": query_id, "sample_id": row["sample_id"],
                "question": row["question"], "company": row["company"],
                "question_type": row["question_type"], "program_type": row["program_type"],
                "evidence_level": level,
                "gold_evidence": json.dumps({
                    "flat": json.loads(compact_evidence(
                        [{"evidence_id": item} for item in gold["flat"]], flat_corpus)),
                    "unitable": json.loads(compact_evidence(
                        [{"evidence_id": item} for item in gold["unitable"]], unitable_corpus)),
                }, ensure_ascii=False, indent=2),
                "flat_top5_evidence": compact_evidence(
                    [{**item, "rank": index} for index, item in enumerate(flat_top, 1)], flat_corpus),
                "unitable_top5_evidence": compact_evidence(
                    [{**item, "rank": index} for index, item in enumerate(unitable_top, 1)],
                    unitable_corpus),
                "flat_gold_rank": int(row["flat_gold_rank"]),
                "unitable_gold_rank": int(row["unitable_gold_rank"]),
                "flat_ndcg@10": float(row["flat_ndcg@10"]),
                "unitable_ndcg@10": float(row["unitable_ndcg@10"]),
                "delta_ndcg": float(row["delta_ndcg"]), "s_teds": float(metric["s_teds"]),
                "valid_html": bool(metric["valid_html"]),
                "cell_count_consistency": bool(metric["cell_count_exact"]),
                "predicted_cells": int(metric["predicted_cells"]),
                "true_cells": int(metric["true_cells"]),
                "predicted_rows": int(metric["predicted_rows"]),
                "true_rows": int(metric["true_rows"]),
                "predicted_cols": int(metric["predicted_cols"]),
                "true_cols": int(metric["true_cols"]),
                "predicted_structure_summary": json.dumps(summary, ensure_ascii=False, indent=2),
                "error_type": "", "manual_note": "", "whether_structure_helpful": "",
            })
        filename, rule = output_specs[group]
        output = args.output_dir / filename
        write_xlsx(output, group, rule, export_rows, args.overwrite, args.seed)
        outputs.append(output)
        LOG.info("Wrote group=%s cases=%d path=%s", group, len(export_rows), output)

    manifest = {
        "schema_version": "1.0", "seed": args.seed, "per_group": args.per_group,
        "selection_uses_model_performance": True,
        "selection_is_deterministic_and_preregistered": True,
        "automatic_error_conclusion": False,
        "selection_rules": {name: rule for name, (_, rule) in output_specs.items()},
        "inputs": [{"path": str(path), "bytes": path.stat().st_size, "sha256": sha256(path)}
                   for path in (args.pairwise, args.flat_retrieval, args.unitable_retrieval,
                                args.flat_corpus, args.unitable_corpus, args.structure_metrics,
                                args.structure_predictions)],
        "outputs": [{"path": str(path), "rows": args.per_group,
                     "bytes": path.stat().st_size, "sha256": sha256(path)} for path in outputs],
    }
    if args.sha_output.exists() and not args.overwrite:
        raise FileExistsError(f"Refusing to overwrite {args.sha_output}; use --overwrite")
    temporary = args.sha_output.with_suffix(args.sha_output.suffix + ".tmp")
    args.sha_output.parent.mkdir(parents=True, exist_ok=True)
    temporary.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
                         encoding="utf-8")
    temporary.replace(args.sha_output)
    LOG.info("Manual-case export complete manifest=%s", args.sha_output)
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        LOG.exception("Manual-case export failed")
        sys.exit(2)
